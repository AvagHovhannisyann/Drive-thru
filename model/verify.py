"""
Recalculates the workbook in LibreOffice (headless) and checks it against the Python engine.
Run: python3 model/verify.py   (exit code 1 on mismatch)
"""
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

import engine

ROOT = Path(__file__).resolve().parent
XL = ROOT / "DriveThru_Financial_Model.xlsx"


def recalc(path: Path) -> Path:
    tmp = Path(tempfile.mkdtemp())
    src = tmp / "in.xlsx"
    shutil.copy(path, src)
    out = tmp / "out"
    out.mkdir()
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out), str(src)],
                   check=True, capture_output=True, timeout=240)
    return out / "in.xlsx"


def main():
    a = engine.load_assumptions()
    calc = recalc(XL)
    wb = load_workbook(calc, data_only=True)
    an = wb["Annual"]
    # find scenario blocks by label
    blocks = {}
    for row in an.iter_rows(min_col=1, max_col=1):
        v = row[0].value
        if isinstance(v, str) and v.endswith("SCENARIO"):
            blocks[v.split()[0].lower()] = row[0].row
    ok = True
    report = {}
    for s, r0 in blocks.items():
        res = engine.run(a, s)
        labels = {}
        for rr in range(r0 + 1, r0 + 80):
            lab = an.cell(row=rr, column=1).value
            if isinstance(lab, str) and lab.endswith("SCENARIO"):
                break
            if lab:
                labels[lab] = rr
        checks = [("Net sales (excl. VAT)", "net"), ("EBITDA", "ebitda"), ("Net income", "net_income"), ("Free cash flow", "fcf"), ("Profit tax", "tax")]
        report[s] = {}
        for lab, key in checks:
            for y in range(5):
                xl = an.cell(row=labels[lab], column=3 + y).value
                py = res["years"][y][key]
                diff = abs((xl or 0) - py)
                tol = max(1.0, abs(py) * 1e-6)
                if diff > tol:
                    ok = False
                    print(f"MISMATCH {s} {lab} Y{y + 1}: excel={xl} python={py}")
            report[s][lab] = [an.cell(row=labels[lab], column=3 + y).value for y in range(5)]
        pb = an.cell(row=labels["Payback (months from opening)"], column=3).value
        report[s]["payback_excel"] = pb
        report[s]["payback_python"] = res["payback_month"]
        report[s]["irr_excel"] = an.cell(row=labels["IRR (5 yrs + residual)"], column=3).value
        report[s]["irr_python"] = res["irr_5y_with_residual"]
        report[s]["npv_excel"] = an.cell(row=labels["NPV @ discount rate (5 yrs + residual)"], column=3).value
        report[s]["npv_python"] = res["npv"]
        if pb != (res["payback_month"] if res["payback_month"] else "> 60"):
            print(f"NOTE payback differs {s}: excel={pb} python={res['payback_month']}")
        if isinstance(report[s]["irr_excel"], float) and res["irr_5y_with_residual"] is not None:
            if abs(report[s]["irr_excel"] - res["irr_5y_with_residual"]) > 1e-4:
                ok = False
                print(f"MISMATCH IRR {s}: {report[s]['irr_excel']} vs {res['irr_5y_with_residual']}")
        if abs((report[s]["npv_excel"] or 0) - res["npv"]) > max(1.0, abs(res["npv"]) * 1e-6):
            ok = False
            print(f"MISMATCH NPV {s}: {report[s]['npv_excel']} vs {res['npv']}")
    be = wb["BreakEven"]
    lab = {be.cell(row=r, column=1).value: r for r in range(1, 40) if be.cell(row=r, column=1).value}
    for j, s in enumerate(["conservative", "base", "optimistic"]):
        res = engine.run(a, s)
        for key, eng_key in (("Break-even transactions/day (EBITDA = 0)", "tx_per_day_ebitda_breakeven"),
                             ("Break-even passing vehicles/day (two-way)", "vehicles_per_day_breakeven"),
                             ("Transactions/day for 4-year cash payback", "tx_per_day_4yr_payback")):
            xl = be.cell(row=lab[key], column=3 + j).value
            py = res["breakeven"][eng_key]
            if abs((xl or 0) - py) > max(0.05, abs(py) * 1e-4):
                ok = False
                print(f"MISMATCH breakeven {s} {key}: excel={xl} python={py}")
            report.setdefault(s, {})[key] = (xl, py)
    (ROOT / "verify_report.json").write_text(json.dumps(report, indent=1, default=str))
    print("VERIFY", "PASS" if ok else "FAIL")
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
